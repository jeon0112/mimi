"""정부 AI/AX 지원사업 자동 수집 → AI 필터 → 이메일 발송

GitHub Actions / 로컬 양용. 매일 실행되어 전 부처의 AI/AX 지원사업·공모를
수집하고 (주)미미에 적합한 공고를 이메일로 정리해 발송한다.
"""

import os
import sys
import smtplib

if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

IS_GITHUB = os.getenv("GITHUB_ACTIONS") == "true"

if IS_GITHUB:
    OUTPUT_DIR = "output"
else:
    OUTPUT_DIR = os.path.join(os.path.expanduser("~"), "Desktop", "정부지원사업결과")

os.makedirs(OUTPUT_DIR, exist_ok=True)

LOG_FILE = os.path.join(OUTPUT_DIR, f"gov_ai_log_{datetime.now().strftime('%Y%m%d')}.txt")


def log(msg: str) -> None:
    line = f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] {msg}"
    print(line, flush=True)
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(line + "\n")


def to_excel(records: list, path: str, detailed: dict = None) -> str:
    """추천/보류 공고 → 엑셀"""
    detailed = detailed or {}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AI_AX_지원사업"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    border = Border(left=Side(style="thin"), right=Side(style="thin"),
                    top=Side(style="thin"), bottom=Side(style="thin"))
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    headers = ["번호", "구분", "D-DAY", "마감일", "지역제한", "공고명", "소관부처", "수행기관",
               "지원분야", "신청기간", "출처", "AI점수", "추천이유", "적합성", "공고URL"]
    widths = [5, 8, 8, 12, 12, 42, 18, 18, 14, 20, 12, 8, 36, 8, 50]

    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 25

    for i, r in enumerate(records, 1):
        e = r.get("평가", {})
        score = e.get("점수", 0)
        status = e.get("추천여부", "")
        rid = r.get("공고ID", "") or r.get("공고명", "")
        fit = detailed.get(rid, {}).get("적합성", "")

        d = r.get("D-DAY")
        dtxt = f"D-{d}" if isinstance(d, int) else "미상"
        row = [
            i, status, dtxt, r.get("마감일", ""), e.get("지역제한", ""),
            r.get("공고명", ""), r.get("소관부처", ""), r.get("수행기관", ""),
            r.get("지원분야", ""), r.get("신청기간", ""), r.get("출처", ""),
            score, e.get("이유", ""), fit, r.get("공고URL", ""),
        ]
        # 마감 임박은 색으로 먼저 보이게 한다.
        if isinstance(d, int) and d <= 7:
            fill = PatternFill("solid", fgColor="FFCDD2")
        elif status == "추천":
            fill = PatternFill("solid", fgColor="E8F5E9")
        else:
            fill = PatternFill("solid", fgColor="FFFDE7")
        ridx = i + 1
        for col, val in enumerate(row, 1):
            cell = ws.cell(row=ridx, column=col, value=val)
            cell.border = border
            cell.fill = fill
            cell.alignment = center if col in [1, 2, 3, 4, 5, 10, 11, 12, 14] else left
        ws.row_dimensions[ridx].height = 38

    ws.cell(row=len(records) + 3, column=1,
            value=f"총 {len(records)}건 | 생성: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    wb.save(path)
    return path


def main():
    log("=" * 50)
    log("정부 AI/AX 지원사업 수집 시작")
    log(f"환경: {'GitHub Actions' if IS_GITHUB else '로컬 PC'}")

    try:
        import gov_ai_collector
        from gov_ai_collector import collect_gov_ai, save_to_json
        from gov_ai_filter import quick_filter, ai_evaluate, ai_detailed_check

        # 1. 수집 (등록일 기준 최근 N일)
        days = 2 if IS_GITHUB else 7
        log("공고 수집 중...")
        records = collect_gov_ai(days=days)
        errors = list(gov_ai_collector.LAST_ERRORS)
        log(f"수집 완료: {len(records)}건")

        # "0건"은 두 가지 상태를 덮는다. 신규 공고가 없는 것과 소스가 죽은 것은 다르다.
        # 소스가 죽었으면 워크플로를 빨간불로 끝내서 눈에 띄게 만든다.
        if errors:
            log("수집 소스 실패:")
            for m in errors:
                log(f"  - {m}")

        if not records:
            if errors:
                log("공고 0건 - 그러나 소스가 실패했다. 공고가 없다고 단정하지 않는다.")
            else:
                log("수집된 공고 없음(소스는 정상). 이메일 발송 후 종료.")
            send_email([], [], None, source_errors=errors)
            if errors:
                sys.exit(1)
            return

        # 2. 사전 필터
        filtered = quick_filter(records)
        log(f"사전 필터 후: {len(filtered)}건")

        # 3. AI 적격 판단
        log("AI 적격 판단 중...")
        evaluations = ai_evaluate(filtered)

        recommended, held = [], []
        for r in filtered:
            e = next((x for x in evaluations if x.get("공고ID") == r.get("공고ID")), None)
            if not e:
                continue
            status = e.get("추천여부", "")
            if status == "추천":
                recommended.append({**r, "평가": e})
            elif status == "보류":
                held.append({**r, "평가": e})

        log(f"추천: {len(recommended)}건 / 보류: {len(held)}건")

        date_str = datetime.now().strftime("%Y%m%d")
        xlsx_path = None
        detailed_checks = {}

        if recommended or held:
            json_path = os.path.join(OUTPUT_DIR, f"gov_ai_{date_str}.json")
            save_to_json(recommended + held, json_path)
            log(f"JSON 저장: {json_path}")

            # 추천 공고 심층 분석
            if recommended:
                log("추천 공고 심층 분석 중...")
                try:
                    detailed_checks = ai_detailed_check(recommended)
                except Exception as e:
                    log(f"심층 분석 오류: {e}")

            xlsx_path = os.path.join(OUTPUT_DIR, f"gov_ai_{date_str}.xlsx")
            to_excel(recommended + held, xlsx_path, detailed_checks)
            log(f"엑셀 저장: {xlsx_path}")

        send_email(recommended, held, xlsx_path, detailed_checks, source_errors=errors)

        log("\n[추천 공고 요약]")
        for r in recommended:
            e = r.get("평가", {})
            d = r.get("D-DAY")
            dtxt = f"D-{d}" if isinstance(d, int) and d >= 0 else "마감일미상"
            region = e.get("지역제한", "")
            rtxt = f" | 지역:{region}" if region and region != "전국" else ""
            log(f"  ✅ [{dtxt}] {r['공고명']} | {r.get('소관부처','')}{rtxt} | {e.get('점수',0)}점")
            log(f"     {r.get('공고URL','')}")

        # 소스 하나라도 죽었으면 결과를 보냈더라도 실패로 끝낸다.
        if errors:
            log("일부 수집 소스가 실패했습니다. 워크플로를 실패로 종료합니다.")
            sys.exit(1)

    except Exception as e:
        log(f"오류: {e}")
        import traceback
        log(traceback.format_exc())
        sys.exit(1)


def send_email(recommended: list, held: list, xlsx_path: str, detailed_checks: dict = None,
               source_errors: list = None) -> None:
    gmail_user = os.getenv("GMAIL_USER", "jjk0112@gmail.com")
    gmail_password = os.getenv("GMAIL_PASSWORD", "")
    to_email = os.getenv("NOTIFY_EMAIL", "jjksp112@naver.com")
    detailed_checks = detailed_checks or {}

    if not gmail_password:
        log("이메일 비밀번호 미설정, 발송 건너뜀")
        return

    source_errors = source_errors or []
    date_str = datetime.now().strftime("%Y년 %m월 %d일")

    def _d(r):
        v = r.get("D-DAY")
        return v if isinstance(v, int) else None

    urgent = [r for r in (recommended + held) if _d(r) is not None and _d(r) <= 7]
    prefix = "[⚠수집실패] " if source_errors else ("[🔥마감임박] " if urgent else "")
    subject = (f"{prefix}[정부 AI/AX 지원사업] {date_str} "
               f"추천 {len(recommended)}건 / 보류 {len(held)}건")

    lines = [
        "안녕하세요, 오늘의 정부 부처 AI/AX 지원사업·공모 결과입니다.\n",
        f"🤖 추천: {len(recommended)}건 | 보류: {len(held)}건",
        "출처: 기업마당(전 부처 통합) + K-Startup\n",
    ]

    # 소스가 죽었으면 맨 위에 말한다. "공고가 없다"로 읽히면 안 된다.
    if source_errors:
        lines.append("\n🚨 [수집 소스 실패] 아래 소스에서 공고를 가져오지 못했습니다.")
        lines.append("   이 메일의 '0건'은 '공고가 없다'는 뜻이 아닙니다.")
        for m in source_errors:
            lines.append(f"   - {m}")
        lines.append("")

    if urgent:
        lines.append("\n🔥 [마감 임박 - D-7 이내]\n")
        for r in sorted(urgent, key=lambda x: _d(x)):
            lines.append(f"• D-{_d(r)} ({r.get('마감일','')}) {r['공고명']}")
            lines.append(f"  {r.get('공고URL','')}")
        lines.append("")

    if recommended:
        lines.append("\n✅ [추천 지원사업]\n")
        for r in recommended:
            e = r.get("평가", {})
            rid = r.get("공고ID", "") or r.get("공고명", "")
            d = _d(r)
            dtxt = f"D-{d}" if d is not None else "마감일 미상"
            lines.append(f"• [{dtxt}] {r['공고명']}")
            region = e.get("지역제한", "")
            rtxt = f" | 지역: {region}" if region and region != "전국" else ""
            lines.append(f"  소관부처: {r.get('소관부처','')} | 수행기관: {r.get('수행기관','')}{rtxt}")
            lines.append(f"  지원분야: {r.get('지원분야','')} | 신청기간: {r.get('신청기간','')}")
            lines.append(f"  AI점수: {e.get('점수',0)}점 - {e.get('이유','')}")
            chk = detailed_checks.get(rid, {})
            if chk:
                lines.append(f"  🔎 적합성: {chk.get('적합성','?')}")
                if chk.get("핵심요약"):
                    lines.append(f"  📌 요약: {chk['핵심요약']}")
                for item in chk.get("체크리스트", [])[:4]:
                    lines.append(f"    {item.get('상태','')} {item.get('항목','')}: {item.get('설명','')}")
                if chk.get("주의사항"):
                    lines.append(f"  ⚠️  주의: {' / '.join(chk['주의사항'][:2])}")
                if chk.get("신청전략"):
                    lines.append(f"  💡 전략: {chk['신청전략']}")
            lines.append(f"  URL: {r.get('공고URL','')}\n")

    if held:
        lines.append("\n⏸ [보류 지원사업]\n")
        for r in held:
            e = r.get("평가", {})
            region = e.get("지역제한", "")
            rtxt = f" [지역: {region}]" if region and region != "전국" else ""
            lines.append(f"• {r['공고명']} ({r.get('소관부처','')}){rtxt} - {e.get('이유','')}")
            lines.append(f"  {r.get('공고URL','')}")

    if not recommended and not held:
        if source_errors:
            lines.append("\n결과 없음 - 다만 위 소스 실패 때문이므로 공고가 없다고 단정할 수 없습니다.")
        else:
            lines.append("\n오늘은 해당하는 AI/AX 지원사업 공고가 없습니다. (수집 소스는 정상)")

    lines.append("\n─────────────────────────")
    lines.append("📚 전체 지원사업 사이트 모음: 저장소의 GOV_AI_SOURCES.md 참고")

    body = "\n".join(lines)

    msg = MIMEMultipart()
    msg["From"] = gmail_user
    msg["To"] = to_email
    msg["Subject"] = subject
    msg.attach(MIMEText(body, "plain", "utf-8"))

    if xlsx_path and os.path.exists(xlsx_path):
        with open(xlsx_path, "rb") as f:
            part = MIMEBase("application", "octet-stream")
            part.set_payload(f.read())
        encoders.encode_base64(part)
        part.add_header("Content-Disposition", f"attachment; filename={os.path.basename(xlsx_path)}")
        msg.attach(part)

    try:
        with smtplib.SMTP_SSL("smtp.gmail.com", 465) as server:
            server.login(gmail_user, gmail_password)
            server.sendmail(gmail_user, to_email, msg.as_string())
        log(f"이메일 발송 완료 → {to_email}")
    except Exception as e:
        log(f"이메일 발송 실패: {e}")


if __name__ == "__main__":
    main()
